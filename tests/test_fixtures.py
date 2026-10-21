import openpyxl


def _ws(path, name):
    return openpyxl.load_workbook(path)[name]


def test_rate_sheet_has_expected_regions(rate_sheet):
    assert openpyxl.load_workbook(rate_sheet).sheetnames == ["Italy", "Greece"]


def test_rate_sheet_italy_has_tiered_rows(rate_sheet):
    ws = _ws(rate_sheet, "Italy")
    assert ws["A4"].value == "Uffizi Gallery — Private Guided Tour"
    assert ws["D4"].value == 285
    assert ws["D5"].value == 240


def test_wide_sheet_puts_pax_bands_in_columns(wide_rate_sheet):
    ws = _ws(wide_rate_sheet, "Japan")
    assert ws["B3"].value == "1-4 Pax"
    assert ws["C3"].value == "5-10 Pax"
    assert ws["A4"].value == "Tea Ceremony — Gion"
    assert ws["B4"].value == 120 and ws["C4"].value == 95


def test_stacked_sheet_has_season_rows_and_usd(stacked_rate_sheet):
    ws = _ws(stacked_rate_sheet, "SouthAfrica")
    assert ws["B3"].value == "Season"
    assert ws["A4"].value == ws["A5"].value  # same experience, two seasons
    assert ws["B4"].value == "Green" and ws["B5"].value == "Peak"
    assert ws["H4"].value == "USD"


def test_multiheader_sheet_has_a_merged_category_band(multiheader_sheet):
    ws = _ws(multiheader_sheet, "Portugal")
    assert "A3:C3" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A3"].value == "Group Size Pricing"
    assert ws["A4"].value == "Experience"      # the real header row
    assert ws["A5"].value == "Douro Valley Wine Tour"


def test_nasty_sheet_has_text_rate_and_blank_row(nasty_sheet):
    ws = _ws(nasty_sheet, "Odd")
    assert ws["C3"].value == "€285"
    assert ws["A5"].value is None


def test_itinerary_has_empty_cost_column(itinerary):
    ws = _ws(itinerary, "Itinerary")
    assert ws["A7"].value == "Experience"
    assert ws["E8"].value is None


def test_day_grouped_itinerary_interleaves_headings(day_grouped_itinerary):
    ws = _ws(day_grouped_itinerary, "Itinerary")
    assert ws["A6"].value == "DAY 1 — FLORENCE"
    assert ws["C6"].value is None            # headings carry no pax
    assert ws["A7"].value == "Uffizi Gallery private tour"
    assert ws["C7"].value == 4


def test_minimal_itinerary_has_no_date_or_time_column(minimal_itinerary):
    ws = _ws(minimal_itinerary, "Itinerary")
    headers = [ws.cell(row=1, column=c).value for c in range(1, 4)]
    assert headers == ["Experience", "Pax", "Cost"]


def test_wide_itinerary_states_pax_once_in_a_header_cell(wide_itinerary):
    ws = _ws(wide_itinerary, "Itinerary")
    assert ws["B3"].value == 4
    assert ws["A5"].value == "Experience"
    headers = [ws.cell(row=5, column=c).value for c in range(1, 5)]
    assert "Pax" not in headers
