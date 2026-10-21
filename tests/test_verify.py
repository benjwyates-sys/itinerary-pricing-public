from decimal import Decimal

import openpyxl

from pricing.catalog import Catalog, CatalogItem, fingerprint
from pricing.itinerary_profile import propose_template
from pricing.match import Aliases
from pricing.verify import refresh, verify_itinerary
from pricing.workbook import Workbook


def _catalog():
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-FLR-0032", name="Uffizi Gallery — Private Guided Tour",
        name_cell="Italy!A4", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "tiered_by_pax", "tiers": [
            {"max_pax": 4, "cell": "Italy!D4", "last_seen": "285"},
            {"max_pax": 10, "cell": "Italy!D5", "last_seen": "240"}]},
        modifiers=[{"type": "day_part", "when": ["PM"], "op": "add_per_unit",
                    "cell": "Italy!F4", "last_seen": "40"}],
        source_rows=[4, 5],
        row_fingerprint=fingerprint("Uffizi Gallery — Private Guided Tour")))
    return catalog


def _hand_priced(itinerary, tmp_path, value):
    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["E8"] = value
    path = tmp_path / "hand.xlsx"
    raw.save(path)
    return path


def _hand_priced_at(itinerary, tmp_path, cell, value, out_name="hand.xlsx"):
    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"][cell] = value
    path = tmp_path / out_name
    raw.save(path)
    return path


def test_matching_hand_price_reports_a_match(itinerary, rate_sheet, tmp_path):
    hand = _hand_priced(itinerary, tmp_path, 1300)
    wb = Workbook(hand)
    diffs, summary = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                      _catalog(), Aliases(), Workbook(rate_sheet))
    row8 = next(d for d in diffs if d.row == 8)
    assert row8.status == "match"
    assert summary["matches"] == 1


def test_mismatch_reports_the_delta(itinerary, rate_sheet, tmp_path):
    hand = _hand_priced(itinerary, tmp_path, 1250)
    wb = Workbook(hand)
    diffs, summary = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                      _catalog(), Aliases(), Workbook(rate_sheet))
    row8 = next(d for d in diffs if d.row == 8)
    assert row8.status == "mismatch"
    assert row8.delta == Decimal("50.00")
    assert summary["mismatches"] == 1


def test_refresh_reports_price_changes(rate_sheet, tmp_path):
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["D4"] = 310
    changed = tmp_path / "changed.xlsx"
    raw.save(changed)

    result = refresh(_catalog(), Workbook(changed))
    assert result["checked"] == 1
    assert result["price_changes"][0]["now"] == "310"
    assert result["broken_items"] == []


def test_refresh_reports_broken_items(rate_sheet, tmp_path):
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"].insert_rows(4)
    raw["Italy"]["A4"] = "Colosseum — Underground Tour"
    shifted = tmp_path / "shifted.xlsx"
    raw.save(shifted)

    result = refresh(_catalog(), Workbook(shifted))
    assert result["broken_items"][0]["id"] == "IT-FLR-0032"


# --- match_confidence: carrying the forced-match signal through the diff ---
#
# verify_itinerary MUST force a resolution for every line (see the module
# docstring in pricing/verify.py) or it could never re-cost the very lines
# it exists to check. But a forced guess and a confirmed match are not the
# same kind of evidence, so every LineDiff now also records item_id (which
# catalog item was actually priced) and match_confidence: "auto_accept"
# (the top candidate really was confident -- a mismatch means the RATE is
# wrong), "top_guess" (forced onto an unconfident top candidate -- a
# mismatch means CHECK THE MATCH FIRST), or "none" (no candidate existed at
# all -- nothing was priced).


def _catamaran_item():
    """Flat-rate item whose name is deliberately NOT similar to the
    itinerary phrasing after normalisation, so the only way it can
    auto_accept is via a saved alias."""
    return CatalogItem(
        id="GR-CAT-0001", name="Catamaran Cruise — PM (beverages incl.)",
        name_cell="Greece!A4", region="Greece", currency="EUR", unit="per_person",
        pricing={"pattern": "flat", "cell": "Greece!D4", "last_seen": "165"},
        source_rows=[4], row_fingerprint="")


def test_saved_alias_match_is_auto_accept(itinerary, rate_sheet, tmp_path):
    catalog = Catalog()
    catalog.upsert(_catamaran_item())
    aliases = Aliases()
    aliases.add("Sunset catamaran w/ open bar", "GR-CAT-0001")

    # Row 10: "Sunset catamaran w/ open bar", pax 4, PM -> 165 * 4 = 660.00
    hand = _hand_priced_at(itinerary, tmp_path, "E10", 660)
    wb = Workbook(hand)
    diffs, _ = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                catalog, aliases, Workbook(rate_sheet))
    row10 = next(d for d in diffs if d.row == 10)
    assert row10.status == "match"
    assert row10.match_confidence == "auto_accept"
    assert row10.item_id == "GR-CAT-0001"


def test_saved_alias_rate_change_is_still_auto_accept_but_mismatches(itinerary, rate_sheet, tmp_path):
    catalog = Catalog()
    catalog.upsert(_catamaran_item())
    aliases = Aliases()
    aliases.add("Sunset catamaran w/ open bar", "GR-CAT-0001")

    raw = openpyxl.load_workbook(rate_sheet)
    raw["Greece"]["D4"] = 200
    changed = tmp_path / "changed_rates.xlsx"
    raw.save(changed)

    # Hand price still reflects the old rate: 165 * 4 = 660, but the sheet
    # now reads 200 * 4 = 800 -- a genuine RATE change, not a bad match.
    hand = _hand_priced_at(itinerary, tmp_path, "E10", 660)
    wb = Workbook(hand)
    diffs, _ = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                catalog, aliases, Workbook(changed))
    row10 = next(d for d in diffs if d.row == 10)
    assert row10.status == "mismatch"
    assert row10.match_confidence == "auto_accept"
    assert row10.delta == Decimal("140.00")


def test_identical_after_normalisation_is_auto_accept(itinerary, rate_sheet, tmp_path):
    # "Vatican Museums skip the line" and "Vatican Museums — Skip the Line"
    # differ only in case/punctuation -- identical after normalisation, so
    # this auto-accepts with no alias involved at all.
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-VAT-0001", name="Vatican Museums — Skip the Line",
        name_cell="Italy!A6", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "flat", "cell": "Italy!D6", "last_seen": "95"},
        source_rows=[6], row_fingerprint=""))

    # Row 9: "Vatican Museums skip the line", pax 4, AM -> 95 * 4 = 380.00
    hand = _hand_priced_at(itinerary, tmp_path, "E9", 380)
    wb = Workbook(hand)
    diffs, _ = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                catalog, Aliases(), Workbook(rate_sheet))
    row9 = next(d for d in diffs if d.row == 9)
    assert row9.status == "match"
    assert row9.match_confidence == "auto_accept"
    assert row9.item_id == "IT-VAT-0001"


def test_tied_decoy_items_are_top_guess_and_item_id_is_visible(itinerary, rate_sheet, tmp_path):
    # Two catalog items whose names are BOTH identical, after normalisation,
    # to the itinerary line's phrasing ("Uffizi Gallery private tour"). This
    # is the genuine tie pricing/match.py guards against: candidates() sets
    # auto_accept back to False on both when more than one item comes back
    # identical to the query, so neither can auto-accept -- but verify still
    # has to force one to price the line. Ties break on item_id ascending,
    # so "IT-DECOY-A" prices ahead of "IT-DECOY-B".
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-DECOY-A", name="Uffizi Gallery private tour",
        name_cell="Italy!A6", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "flat", "cell": "Italy!D6", "last_seen": "95"},
        source_rows=[6], row_fingerprint=""))
    catalog.upsert(CatalogItem(
        id="IT-DECOY-B", name="Uffizi Gallery private tour",
        name_cell="Italy!A7", region="Italy", currency="EUR", unit="per_group",
        pricing={"pattern": "flat", "cell": "Italy!D7", "last_seen": "110"},
        source_rows=[7], row_fingerprint=""))

    # Row 8 was hand-priced at 1300 (the real Uffizi rate), but neither
    # decoy prices anywhere near that -- the mismatch this produces is a
    # bad-MATCH mismatch, not a bad-RATE one, and match_confidence says so.
    hand = _hand_priced(itinerary, tmp_path, 1300)
    wb = Workbook(hand)
    diffs, _ = verify_itinerary(hand, propose_template(wb, "Itinerary"),
                                catalog, Aliases(), Workbook(rate_sheet))
    row8 = next(d for d in diffs if d.row == 8)
    assert row8.match_confidence == "top_guess"
    assert row8.item_id == "IT-DECOY-A"  # visible: which wrong item was priced
    assert row8.computed == Decimal("380.00")  # 95 * 4 pax -- the wrong item's rate
    assert row8.status == "mismatch"


def test_unrelated_line_against_single_item_catalog_is_top_guess(itinerary, rate_sheet):
    # Rows 9 ("Vatican Museums skip the line") and 10 ("Sunset catamaran w/
    # open bar") have nothing to do with the only catalog item (Uffizi
    # Gallery). Forcing still prices them against it -- but the result must
    # not look confident.
    wb = Workbook(itinerary)
    diffs, _ = verify_itinerary(itinerary, propose_template(wb, "Itinerary"),
                                _catalog(), Aliases(), Workbook(rate_sheet))
    for row in (9, 10):
        d = next(x for x in diffs if x.row == row)
        assert d.match_confidence == "top_guess"
        assert d.item_id == "IT-FLR-0032"


def test_empty_catalog_reports_none_confidence_and_not_priced(itinerary, rate_sheet):
    wb = Workbook(itinerary)
    diffs, _ = verify_itinerary(itinerary, propose_template(wb, "Itinerary"),
                                Catalog(), Aliases(), Workbook(rate_sheet))
    row8 = next(d for d in diffs if d.row == 8)
    assert row8.match_confidence == "none"
    assert row8.item_id is None
    assert row8.computed is None
    assert row8.status == "not_priced"


def test_summary_top_guess_count_reconciles_with_the_lines(itinerary, rate_sheet):
    # Single-item catalog against all three itinerary lines: only "Uffizi
    # Gallery private tour" is even plausibly related, and it still isn't
    # identical after normalisation, so all three lines are forced guesses.
    wb = Workbook(itinerary)
    diffs, summary = verify_itinerary(itinerary, propose_template(wb, "Itinerary"),
                                      _catalog(), Aliases(), Workbook(rate_sheet))
    top_guess_lines = [d for d in diffs if d.match_confidence == "top_guess"]
    assert summary["top_guesses"] == len(top_guess_lines)
    assert summary["top_guesses"] == 3


# --- item_id must never be None while match_confidence claims a match was
# found (Task 14 review, fixed in Task 16) ---
#
# price_lines() short-circuits BEFORE it ever sets PricedLine.item_id when a
# line's pax cell is unparseable (see pricing/cost.py's pax_kind handling):
# the line becomes needs_quote/not_priced without ever reaching the candidate
# resolution step. But verify_itinerary computes forced_resolutions and
# match_confidence for every line up front, independent of whether pax could
# be read. Reading item_id back off the PricedLine (which never got one)
# instead of off forced_resolutions (which did) left LineDiff.item_id == None
# next to match_confidence == "top_guess"/"auto_accept" -- contradicting
# LineDiff.item_id's own docstring ("None only when nothing was priced at
# all... match_confidence == 'none'").


def test_unparseable_pax_still_reports_the_forced_item_id(itinerary, rate_sheet, tmp_path):
    # The template is profiled from the itinerary's UNEDITED layout, same as
    # real usage (it's confirmed once in Phase 0.5, then reused) -- deriving
    # it from the edited file instead would let the header-row heuristic
    # itself react to C8 losing its only numeric cell, which is a confound
    # this test has nothing to do with.
    template = propose_template(Workbook(itinerary), "Itinerary")

    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["C8"] = "4-6"  # unparseable party size, not a bare integer
    path = tmp_path / "ambiguous_pax.xlsx"
    raw.save(path)

    diffs, _ = verify_itinerary(path, template,
                                _catalog(), Aliases(), Workbook(rate_sheet))
    row8 = next(d for d in diffs if d.row == 8)
    assert row8.status == "not_priced"  # pax couldn't be read, so nothing was computed
    assert row8.match_confidence != "none"  # a candidate WAS found for this line
    assert row8.item_id == "IT-FLR-0032"  # ...and it must be visible here, not None
