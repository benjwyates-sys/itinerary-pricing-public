from decimal import Decimal

import openpyxl
import pytest

from pricing.catalog import (
    Catalog,
    CatalogItem,
    FingerprintMismatch,
    compute_structure_fingerprint,
    fingerprint,
    resolve_rates,
)
from pricing.workbook import Workbook

UFFIZI = CatalogItem(
    id="IT-FLR-0032",
    name="Uffizi Gallery — Private Guided Tour",
    name_cell="Italy!A4",
    region="Italy",
    currency="EUR",
    unit="per_person",
    pricing={
        "pattern": "tiered_by_pax",
        "tiers": [
            {"max_pax": 4, "cell": "Italy!D4", "last_seen": "285"},
            {"max_pax": 10, "cell": "Italy!D5", "last_seen": "240"},
        ],
    },
    modifiers=[
        {"type": "day_part", "when": ["PM"], "op": "add_per_unit",
         "cell": "Italy!F4", "last_seen": "40"}
    ],
    source_rows=[4, 5],
    row_fingerprint=fingerprint("Uffizi Gallery — Private Guided Tour"),
    confirmed_at="2026-09-03",
)


def test_catalog_roundtrips(tmp_path):
    catalog = Catalog()
    catalog.upsert(UFFIZI)
    path = tmp_path / "catalog.json"
    catalog.save(path)
    assert Catalog.load(path).get("IT-FLR-0032").name == UFFIZI.name


def test_upsert_replaces_rather_than_duplicates(tmp_path):
    catalog = Catalog()
    catalog.upsert(UFFIZI)
    catalog.upsert(UFFIZI)
    assert len(catalog.items) == 1


def test_load_missing_file_gives_empty_catalog(tmp_path):
    assert Catalog.load(tmp_path / "nope.json").items == []


def test_resolves_live_rates(rate_sheet):
    rates, changes = resolve_rates(UFFIZI, Workbook(rate_sheet))
    assert rates["Italy!D4"] == Decimal("285")
    assert rates["Italy!F4"] == Decimal("40")
    assert changes == []


def test_detects_a_price_change(rate_sheet, tmp_path):
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["D4"] = 310
    changed = tmp_path / "changed.xlsx"
    raw.save(changed)

    rates, changes = resolve_rates(UFFIZI, Workbook(changed))
    assert rates["Italy!D4"] == Decimal("310")
    assert len(changes) == 1
    assert changes[0].last_seen == Decimal("285")
    assert changes[0].current == Decimal("310")


def test_shifted_rows_refuse_to_price(rate_sheet, tmp_path):
    """The silent-and-expensive failure this design exists to prevent."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"].insert_rows(4)
    raw["Italy"]["A4"] = "Colosseum — Underground Tour"
    shifted = tmp_path / "shifted.xlsx"
    raw.save(shifted)

    with pytest.raises(FingerprintMismatch) as exc:
        resolve_rates(UFFIZI, Workbook(shifted))
    assert exc.value.item_id == "IT-FLR-0032"
    assert "map" in str(exc.value).lower()


def test_fingerprint_ignores_whitespace_and_case():
    assert fingerprint("  Uffizi  Gallery ") == fingerprint("uffizi gallery")


# Additional tests for structural fingerprint (drift detection).


@pytest.fixture
def uffizi_with_structure(rate_sheet):
    """UFFIZI with structure_fingerprint computed from the rate sheet."""
    wb = Workbook(rate_sheet)
    structure_fp = compute_structure_fingerprint(wb, UFFIZI)
    item = CatalogItem(
        id="IT-FLR-0032",
        name="Uffizi Gallery — Private Guided Tour",
        name_cell="Italy!A4",
        region="Italy",
        currency="EUR",
        unit="per_person",
        pricing={
            "pattern": "tiered_by_pax",
            "tiers": [
                {"max_pax": 4, "cell": "Italy!D4", "last_seen": "285"},
                {"max_pax": 10, "cell": "Italy!D5", "last_seen": "240"},
            ],
        },
        modifiers=[
            {"type": "day_part", "when": ["PM"], "op": "add_per_unit",
             "cell": "Italy!F4", "last_seen": "40"}
        ],
        source_rows=[4, 5],
        row_fingerprint=fingerprint("Uffizi Gallery — Private Guided Tour"),
        structure_fingerprint=structure_fp,
        confirmed_at="2026-09-03",
    )
    wb.close()
    return item


def test_column_inserted_before_rate_triggers_structure_mismatch(
    uffizi_with_structure, rate_sheet, tmp_path
):
    """Column inserted before the rate column shifts all cells."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"].insert_cols(4)  # Insert before column D (rate column)
    shifted = tmp_path / "col_inserted.xlsx"
    raw.save(shifted)

    with pytest.raises(FingerprintMismatch) as exc:
        resolve_rates(uffizi_with_structure, Workbook(shifted))
    assert exc.value.item_id == "IT-FLR-0032"
    assert "map" in str(exc.value).lower()


def test_row_deleted_above_item_triggers_structure_mismatch(
    uffizi_with_structure, rate_sheet, tmp_path
):
    """Row deleted above the item shifts rows and breaks structure fingerprint."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"].delete_rows(3)  # Delete row 3, shifting rows 4,5 up to 3,4
    shifted = tmp_path / "row_deleted.xlsx"
    raw.save(shifted)

    with pytest.raises(FingerprintMismatch) as exc:
        resolve_rates(uffizi_with_structure, Workbook(shifted))
    assert exc.value.item_id == "IT-FLR-0032"
    assert "map" in str(exc.value).lower()


def test_legitimate_price_update_does_not_trigger_mismatch(
    uffizi_with_structure, rate_sheet, tmp_path
):
    """Only the rate cell changes; structure fingerprint stays stable, price change is reported."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["D4"] = 310  # Update only the rate cell
    updated = tmp_path / "price_updated.xlsx"
    raw.save(updated)

    # Should not raise FingerprintMismatch; should report PriceChange instead.
    rates, changes = resolve_rates(uffizi_with_structure, Workbook(updated))
    assert rates["Italy!D4"] == Decimal("310")
    assert len(changes) == 1
    assert changes[0].last_seen == Decimal("285")
    assert changes[0].current == Decimal("310")


def test_empty_structure_fingerprint_skips_verification(rate_sheet):
    """Empty structure_fingerprint means verification is skipped (not an error)."""
    # UFFIZI without structure_fingerprint should still work (just unprotected).
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"].insert_cols(4)  # Structural change
    tmp_path_used = rate_sheet.parent / "temp_unprotected.xlsx"
    raw.save(tmp_path_used)

    # Item with NO structure_fingerprint should not raise, just load rates.
    item_no_structure = CatalogItem(
        id="IT-FLR-0032",
        name="Uffizi Gallery — Private Guided Tour",
        name_cell="Italy!A4",
        region="Italy",
        currency="EUR",
        unit="per_person",
        pricing={
            "pattern": "tiered_by_pax",
            "tiers": [
                {"max_pax": 4, "cell": "Italy!D4", "last_seen": "285"},
                {"max_pax": 10, "cell": "Italy!D5", "last_seen": "240"},
            ],
        },
        modifiers=[
            {"type": "day_part", "when": ["PM"], "op": "add_per_unit",
             "cell": "Italy!F4", "last_seen": "40"}
        ],
        source_rows=[4, 5],
        row_fingerprint="",  # No row fingerprint
        structure_fingerprint="",  # No structure fingerprint
        confirmed_at="2026-09-03",
    )

    # Should not raise, even though structure changed.
    # This is the key difference: with empty fingerprints, no verification occurs.
    rates, _ = resolve_rates(item_no_structure, Workbook(tmp_path_used))
    # Note: rates may be empty or wrong due to the column shift, but no exception raised.
    # The point is that empty fingerprint means unprotected, not that it catches the error.

    tmp_path_used.unlink()  # Clean up
