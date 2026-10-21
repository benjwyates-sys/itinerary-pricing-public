from pathlib import Path

import openpyxl
import pytest

from pricing.itinerary_profile import MissingRoleError, propose_template
from pricing.workbook import Workbook


def test_baseline_is_high_confidence(itinerary):
    template = propose_template(Workbook(itinerary), "Itinerary")
    assert template.confidence == "high"
    assert template.roles["cost"] == "E"


def test_day_grouped_itinerary_finds_the_same_roles(day_grouped_itinerary):
    template = propose_template(Workbook(day_grouped_itinerary), "Itinerary")
    assert template.header_row == 5
    assert template.first_data_row == 6
    assert template.roles["pax"] == "C"


def test_minimal_itinerary_has_no_date_or_time_role(minimal_itinerary):
    template = propose_template(Workbook(minimal_itinerary), "Itinerary")
    assert template.roles["name"] == "A"
    assert template.roles["cost"] == "C"
    assert "date" not in template.roles
    assert "day_part" not in template.roles


def test_wide_itinerary_has_no_pax_column_and_says_so(wide_itinerary):
    """Adversarial: pax is stated once in a header cell, not per row."""
    template = propose_template(Workbook(wide_itinerary), "Itinerary")
    assert "pax" not in template.roles
    assert template.confidence == "low"
    assert template.default_pax is None       # the operator supplies it


def test_alternatives_lists_the_rejected_cost_column(tmp_path: Path):
    """Two columns both plausibly hold the cost role ('Net Rate' matches
    'net', 'Total Amount' matches 'total'/'amount'). 'net' has keyword
    priority so 'Net Rate' is chosen; 'Total Amount' must still be
    surfaced as an alternative -- and the chosen column must NOT appear
    in its own alternatives list."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Itinerary"
    ws["A7"] = "Experience"
    ws["B7"] = "Date"
    ws["C7"] = "Pax"
    ws["D7"] = "Net Rate"
    ws["E7"] = "Total Amount"
    ws["A8"] = "Hotel"
    ws["B8"] = "2025-01-01"
    ws["C8"] = 2
    path = tmp_path / "ambiguous_cost.xlsx"
    wb.save(path)

    template = propose_template(Workbook(path), "Itinerary")
    assert isinstance(template.alternatives, dict)
    assert template.roles["cost"] == "D"
    assert "E" in template.alternatives["cost"]
    assert "D" not in template.alternatives["cost"]


def test_alternatives_lists_the_rejected_pax_column(tmp_path: Path):
    """'Pax' and 'Adults' both plausibly hold the pax role. 'pax' has
    keyword priority so 'Pax' is chosen; 'Adults' must still be surfaced
    as an alternative."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Itinerary"
    ws["A7"] = "Experience"
    ws["B7"] = "Date"
    ws["C7"] = "Pax"
    ws["D7"] = "Adults"
    ws["E7"] = "Cost"
    ws["A8"] = "Hotel"
    ws["B8"] = "2025-01-01"
    ws["C8"] = 4
    ws["D8"] = 4
    ws["E8"] = 100
    path = tmp_path / "ambiguous_pax.xlsx"
    wb.save(path)

    template = propose_template(Workbook(path), "Itinerary")
    assert template.roles["pax"] == "C"
    assert "D" in template.alternatives["pax"]
    assert "C" not in template.alternatives["pax"]


def test_missing_cost_column_is_still_an_error(itinerary, tmp_path):
    import openpyxl

    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["E7"] = "Notes"
    stripped = tmp_path / "no_cost.xlsx"
    raw.save(stripped)

    with pytest.raises(MissingRoleError) as exc:
        propose_template(Workbook(stripped), "Itinerary")
    assert "cost" in exc.value.missing
