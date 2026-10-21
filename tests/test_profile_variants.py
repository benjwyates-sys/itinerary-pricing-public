import json
from pathlib import Path

import openpyxl

from pricing.profile import find_header_candidates, profile_sheet, profile_workbook, to_dict
from pricing.workbook import Workbook


def test_wide_sheet_headers_are_captured_verbatim(wide_rate_sheet):
    profile = profile_sheet(Workbook(wide_rate_sheet), "Japan")
    headers = [c.header for c in profile.columns]
    assert "1-4 Pax" in headers and "5-10 Pax" in headers and "11+ Pax" in headers


def test_wide_sheet_pax_band_columns_are_price_bearing(wide_rate_sheet):
    """'1-4 Pax' holds money despite the word 'Pax' in its header."""
    profile = profile_sheet(Workbook(wide_rate_sheet), "Japan")
    band = next(c for c in profile.columns if c.header == "1-4 Pax")
    assert band.price_bearing is True


def test_stacked_sheet_exposes_season_and_currency_values(stacked_rate_sheet):
    profile = profile_sheet(Workbook(stacked_rate_sheet), "SouthAfrica")
    by_header = {c.header: c for c in profile.columns}
    assert sorted(by_header["Season"].sample_values) == ["All Year", "Green", "Peak"]
    assert by_header["Currency"].sample_values == ["USD"]


def test_multiheader_offers_both_candidate_rows(multiheader_sheet):
    """Adversarial: row 3 is a MERGED category band, row 4 is the real header."""
    wb = Workbook(multiheader_sheet)
    max_row, max_col = wb.dimensions("Portugal")
    candidates = find_header_candidates(wb, "Portugal", max_row, max_col)
    assert 3 in candidates and 4 in candidates
    # The real header must outrank the band, not just appear somewhere in the list.
    assert candidates == [4, 3]


def test_multiheader_profile_records_the_candidates(multiheader_sheet):
    profile = profile_sheet(Workbook(multiheader_sheet), "Portugal")
    assert len(profile.header_row_candidates) >= 2
    assert profile.header_row == 4


def _find_candidates(path: Path, sheet: str) -> list[int]:
    wb = Workbook(path)
    max_row, max_col = wb.dimensions(sheet)
    return find_header_candidates(wb, sheet, max_row, max_col)


def test_unmerged_band_still_ranks_the_real_header_first(tmp_path: Path):
    """Adversarial, and NOT covered by the merged-range penalty: the band
    row's cells are typed across adjacent columns without merging them --
    an entirely ordinary way to build a spreadsheet. The only reliable
    signal left is structural: the band sits above ANOTHER header-like
    row, not directly above data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"] = (
        "Group", "Size", "Pricing", "Band", "Table")
    ws["A4"], ws["B4"], ws["C4"] = "Experience", "Rate", "Unit"
    ws["A5"], ws["B5"], ws["C5"] = "Tour A", 100, "per person"
    ws["A6"], ws["B6"], ws["C6"] = "Tour B", 150, "per person"
    path = tmp_path / "unmerged_band.xlsx"
    wb.save(path)

    candidates = _find_candidates(path, "Sheet1")
    assert candidates[0] == 4
    assert candidates == [4, 3]


def test_three_row_header_stack_ranks_the_row_above_data_first(tmp_path: Path):
    """Two layers of category text sit above the real, three-column header;
    only the row directly above numeric data should win."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A3"], ws["B3"], ws["C3"] = "Group Size", "Pricing Tiers", "Notes Column"
    ws["A4"], ws["B4"], ws["C4"] = "Sub Category", "Rate Bands", "Extra Info"
    ws["A5"], ws["B5"], ws["C5"] = "Experience", "Rate", "Unit"
    ws["A6"], ws["B6"], ws["C6"] = "Tour A", 100, "per person"
    path = tmp_path / "three_row_stack.xlsx"
    wb.save(path)

    candidates = _find_candidates(path, "Sheet1")
    assert candidates[0] == 5


def test_band_at_row_one_still_ranks_header_at_row_two_first(tmp_path: Path):
    """No title block above the band this time -- the band is the very
    first row in the sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Group", "Size", "Pricing", "Band"
    ws["A2"], ws["B2"], ws["C2"] = "Experience", "Rate", "Unit"
    ws["A3"], ws["B3"], ws["C3"] = "Tour A", 100, "per person"
    path = tmp_path / "band_row1.xlsx"
    wb.save(path)

    candidates = _find_candidates(path, "Sheet1")
    assert candidates[0] == 2


def test_band_directly_above_data_is_not_penalized(tmp_path: Path):
    """Regression guard for the new penalty: when a band-looking row IS the
    real header -- it sits directly above numeric data, with no separate
    header row beneath it -- it must not be punished for merely looking
    like a band."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A3"], ws["B3"], ws["C3"] = "Group Size Pricing", "Rate", "Unit"
    ws["A4"], ws["B4"], ws["C4"] = "Tour A", 100, "per person"
    ws["A5"], ws["B5"], ws["C5"] = "Tour B", 150, "per person"
    path = tmp_path / "band_with_numbers_beneath.xlsx"
    wb.save(path)

    candidates = _find_candidates(path, "Sheet1")
    assert candidates == [3]


def test_existing_fixtures_do_not_regress(rate_sheet, wide_rate_sheet,
                                          stacked_rate_sheet, nasty_sheet):
    """The new stacked-header penalty must not disturb any layout that was
    already ranking correctly."""
    assert _find_candidates(rate_sheet, "Italy")[0] == 3
    assert _find_candidates(wide_rate_sheet, "Japan")[0] == 3
    assert _find_candidates(stacked_rate_sheet, "SouthAfrica")[0] == 3
    assert _find_candidates(nasty_sheet, "Odd")[0] == 2


def test_no_variant_leaks_a_rate(rate_sheet, wide_rate_sheet, stacked_rate_sheet,
                                 multiheader_sheet, nasty_sheet):
    """The safe-to-share promise must hold for every layout, not just one."""
    secrets = {
        rate_sheet: ["285", "240", "110", "95", "165", "40"],
        wide_rate_sheet: ["120", "95", "80", "140", "85", "70", "60", "15"],
        stacked_rate_sheet: ["180", "240", "310"],
        multiheader_sheet: ["210", "175", "195", "160", "25"],
        nasty_sheet: ["285", "1234", "1.234,50"],
    }
    for path, values in secrets.items():
        blob = json.dumps(to_dict(profile_workbook(Workbook(path))))
        for secret in values:
            assert secret not in blob, f"{path.name} leaked {secret}"
