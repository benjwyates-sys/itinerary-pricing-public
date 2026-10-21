"""Tests for pricing.cost: turning parsed itinerary lines into priced lines.

IMPORTANT: `pricing.match.candidates()` was redesigned after three failed
tuning rounds (see pricing/match.py's module docstring and tests/test_match.py).
A line is only auto-priced without asking when the top candidate's
`auto_accept` flag is True, which happens in exactly two cases: the names
are identical after normalisation, or the query hits a saved alias. It is
NEVER a similarity-score threshold.

"Uffizi Gallery private tour" (the itinerary's phrasing) is deliberately
NOT identical, after normalisation, to "Uffizi Gallery — Private Guided
Tour" (the catalog's phrasing) — that exact pair is asserted to NOT
auto-accept in tests/test_match.py's MUST_NOT_AUTO_ACCEPT corpus. So tests
here that need that line priced without an operator resolution seed a
saved alias for it first, which is one of the two legitimate auto-accept
paths. "Vatican Museums skip the line" happens to be identical to
"Vatican Museums — Skip the Line" after normalisation (case and dash
folding only) and needs no alias.
"""

import itertools
import types
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from pricing.catalog import Catalog, CatalogItem, fingerprint
from pricing.cost import parse_itinerary, price_lines
from pricing.itinerary_profile import propose_template
from pricing.match import Aliases
from pricing.workbook import Workbook


def _catalog():
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-FLR-0032", name="Uffizi Gallery — Private Guided Tour",
        name_cell="Italy!A4", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "tiered_by_pax", "tiers": [
            {"max_pax": 4, "cell": "Italy!D4", "last_seen": "285"},
            {"max_pax": 10, "cell": "Italy!D5", "last_seen": "240"}]},
        modifiers=[{"type": "day_part", "when": ["PM"], "op": "add_per_unit",
                    "cell": "Italy!F4", "last_seen": "40"}],
        source_rows=[4, 5],
        row_fingerprint=fingerprint("Uffizi Gallery — Private Guided Tour")))
    catalog.upsert(CatalogItem(
        id="IT-ROM-0001", name="Vatican Museums — Skip the Line",
        name_cell="Italy!A6", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "flat", "cell": "Italy!D6", "last_seen": "95"},
        source_rows=[6], row_fingerprint=fingerprint("Vatican Museums — Skip the Line")))
    return catalog


def _aliases_with_uffizi():
    """An operator has previously confirmed this exact itinerary phrasing.

    This is the alias-hit auto_accept path, not a score threshold — the
    only other legitimate path is identical-after-normalisation, which is
    what covers "Vatican Museums skip the line" below without any alias.
    """
    aliases = Aliases()
    aliases.add("Uffizi Gallery private tour", "IT-FLR-0032")
    return aliases


def test_parses_lines_from_the_template(itinerary):
    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    assert [l.row for l in lines] == [8, 9, 10]
    assert lines[0].raw_name == "Uffizi Gallery private tour"
    assert lines[0].pax == 4
    assert lines[0].date == date(2026, 6, 15)
    assert lines[0].day_part == "PM"


def test_prices_confident_matches(itinerary, rate_sheet):
    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, changes = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))

    uffizi = next(p for p in priced if p.row == 8)
    assert uffizi.status == "priced"
    assert uffizi.item_id == "IT-FLR-0032"
    assert uffizi.total == Decimal("1300.00")   # (285 + 40 PM) * 4
    assert changes == []

    # "Vatican Museums skip the line" is identical to the catalog's
    # "Vatican Museums — Skip the Line" after normalisation (case/dash
    # folding only) — a genuine auto_accept with no alias needed.
    vatican = next(p for p in priced if p.row == 9)
    assert vatican.status == "priced"
    assert vatican.item_id == "IT-ROM-0001"
    assert vatican.total == Decimal("380.00")   # 95 * 4, no PM supplement


def test_unrecognised_line_is_unmatched_with_candidates(itinerary, rate_sheet):
    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(rate_sheet))

    catamaran = next(p for p in priced if p.row == 10)
    assert catamaran.status == "unmatched"
    assert catamaran.total is None
    assert len(catamaran.candidates) <= 3


def test_operator_resolution_prices_the_line_and_is_honoured(itinerary, rate_sheet):
    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(rate_sheet),
                            resolutions={10: "IT-ROM-0001"})
    catamaran = next(p for p in priced if p.row == 10)
    assert catamaran.status == "priced"
    assert catamaran.item_id == "IT-ROM-0001"


def test_needs_quote_is_reported_not_guessed(itinerary, rate_sheet, tmp_path):
    import openpyxl
    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["C8"] = 14          # above the top tier
    big = tmp_path / "big.xlsx"
    raw.save(big)

    wb = Workbook(big)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))
    line = next(p for p in priced if p.row == 8)
    assert line.status == "needs_quote"
    assert line.total is None
    assert "14" in line.reason


def test_day_headings_are_not_priced(day_grouped_itinerary, rate_sheet):
    """'DAY 1 - FLORENCE' has a name but no pax. It is not a line item."""
    wb = Workbook(day_grouped_itinerary)
    lines, skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    assert [l.row for l in lines] == [7, 9]
    assert all("DAY" not in l.raw_name for l in lines)
    # C3 (code review): headings are still reported as skipped rows -- the
    # heuristic that tells a heading apart from a real line item with a
    # forgotten pax cell cannot actually tell them apart, so both are
    # surfaced rather than silently dropped. It's on the human, not a
    # guess, to say "yes, that one's just a heading."
    assert {s["row"]: s["name"] for s in skipped} == {
        6: "DAY 1 — FLORENCE", 8: "DAY 2 — ROME",
    }


# --- C3: a dropped row is never invisible ----------------------------------
#
# Three real paths land on the same "name present, pax cell blank" heading
# heuristic and used to vanish the row with no trace anywhere: an uncached
# formula in the pax cell, a stale saved template applied to a
# differently-shaped file (covered in tests/test_cli.py), and a genuinely
# blank pax cell on a real line ("same party throughout"). All three now
# show up in parse_itinerary's second return value.

def test_uncached_formula_in_pax_cell_is_reported_not_silently_dropped(
        itinerary, rate_sheet, tmp_path):
    """A pax cell holding a formula that has never been calculated by Excel
    (e.g. built fresh with openpyxl, never opened and saved by Excel
    itself) reads back as None through the data_only view -- identical to
    a genuinely blank cell. Before this fix the Uffizi line just vanished;
    now it's absent from `lines` (never priced with a guessed pax) but
    named in `skipped` instead of disappearing without a trace."""
    # Template derived from the ORIGINAL, unmutated workbook -- same
    # reasoning as _itinerary_with_pax_cell above: a formula string in
    # what is normally a numeric pax cell can otherwise confound
    # find_header_row's own heuristics, which is not what this test is
    # about.
    template = propose_template(Workbook(itinerary), "Itinerary")

    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["C8"] = "=G1"   # uncached formula -- reads as None
    formula_variant = tmp_path / "formula_pax.xlsx"
    raw.save(formula_variant)

    wb = Workbook(formula_variant)
    lines, skipped = parse_itinerary(wb, template)

    assert all(l.row != 8 for l in lines)
    assert any(s["row"] == 8 and "Uffizi" in s["name"] for s in skipped)


def test_blank_pax_on_a_real_line_is_reported_not_silently_dropped(
        itinerary, tmp_path):
    """A genuinely blank pax cell on a real line item ('same party
    throughout' -- never filled in per-row) must be visible, not silently
    treated as a heading and dropped."""
    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"]["C9"] = None
    blank_variant = tmp_path / "blank_pax.xlsx"
    raw.save(blank_variant)

    wb = Workbook(blank_variant)
    template = propose_template(wb, "Itinerary")
    lines, skipped = parse_itinerary(wb, template)

    assert all(l.row != 9 for l in lines)
    assert any(s["row"] == 9 and "Vatican" in s["name"] for s in skipped)


_variant_counter = itertools.count()


def _itinerary_with_pax_cell(itinerary, tmp_path, cell_value, row=8):
    """Return (template, mutated_workbook) for a copy of the `itinerary`
    fixture with the given row's pax cell (column C -- row 8 is 'Uffizi
    Gallery private tour', row 9 is 'Vatican Museums skip the line')
    replaced by an arbitrary value (text, a bare number, or even a date
    object), so we can exercise pax parsing without touching
    tests/fixtures/build.py.

    The template is derived from the ORIGINAL, unmutated workbook: swapping
    a numeric pax cell for text like "4 adults" or "four" makes that row
    look more header-like than the real header row to find_header_row()'s
    heuristics (a profile.py concern, out of scope here), which would
    otherwise misdetect the header row on the mutated copy. The column
    layout itself doesn't change, so reusing the template computed from the
    original file is valid and sidesteps that unrelated heuristic.
    """
    template = propose_template(Workbook(itinerary), "Itinerary")

    raw = openpyxl.load_workbook(itinerary)
    raw["Itinerary"][f"C{row}"] = cell_value
    variant = tmp_path / f"pax_variant_{next(_variant_counter)}.xlsx"
    raw.save(variant)
    return template, Workbook(variant)


# --- Strict pax parsing --------------------------------------------------
#
# Round 2 of this fix: a prior version accepted "the first run of digits
# anywhere at the start of the cell" -- which silently under-charged
# "2 adults + 1 child" as 2, collapsed "4-6" and "4+" to 4, priced "0" as a
# free experience, and turned a stray DATE into a five-figure line. Every
# one of those must now come back as needs_quote instead.
#
# A cell is "parsed" ONLY when it is: an integer-valued number cell; text
# that is exactly a bare run of digits (optionally with one trailing ","
# or "."); or text that is exactly "<digits> <noun>" for a small set of
# party-size nouns -- and the integer is 1-200 inclusive. Everything else
# is unparseable (or, if it IS an unambiguous integer but out of bounds,
# "out_of_range") and must ask, never guess.

PARSED_PAX_CASES = [
    pytest.param(4, id="int-4"),
    pytest.param(4.0, id="float-4.0"),
    pytest.param("4", id="text-4"),
    pytest.param(" 4 ", id="text-padded-4"),
    pytest.param("04", id="text-leading-zero-04"),
    pytest.param("4,", id="text-trailing-comma-4,"),
    pytest.param("4 pax", id="text-4-pax"),
    pytest.param("4 adults", id="text-4-adults"),
    pytest.param("4 people", id="text-4-people"),
    pytest.param("4 guests", id="text-4-guests"),
    pytest.param("4 pp", id="text-4-pp"),
]

NEEDS_QUOTE_PAX_CASES = [
    pytest.param("0", id="zero"),
    pytest.param("-2", id="negative"),
    pytest.param("4.5", id="non-integer"),
    pytest.param("201", id="above-ceiling"),
    pytest.param("99999", id="way-above-ceiling"),
    pytest.param("2 adults + 1 child", id="party-breakdown"),
    pytest.param("4-6", id="range"),
    pytest.param("4+", id="open-count"),
    pytest.param("1x4", id="multiplication"),
    pytest.param("4 (TBC)", id="unconfirmed"),
    pytest.param("Pax: 4", id="label-prefixed"),
    pytest.param("approx 4", id="approximate"),
    pytest.param("four", id="spelled-out"),
    pytest.param("TBC", id="no-number-at-all"),
    pytest.param(date(2026, 6, 15), id="stray-date"),
]


@pytest.mark.parametrize("pax_value", PARSED_PAX_CASES)
def test_pax_parsing_is_strict_about_what_counts_as_unambiguous(
        itinerary, rate_sheet, tmp_path, pax_value):
    """Only a narrow set of unambiguous forms are parsed automatically --
    each of these resolves to a party of 4 and prices normally."""
    template, wb = _itinerary_with_pax_cell(itinerary, tmp_path, pax_value)

    lines, _skipped = parse_itinerary(wb, template)
    uffizi_line = next(l for l in lines if l.row == 8)
    assert uffizi_line.pax == 4
    assert uffizi_line.pax_kind == "parsed"

    priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))
    uffizi = next(p for p in priced if p.row == 8)
    assert uffizi.status == "priced"
    assert uffizi.total == Decimal("1300.00")   # (285 + 40 PM) * 4


def test_pax_parsing_accepts_the_upper_plausibility_bound(itinerary, rate_sheet, tmp_path):
    """200 is the top of the accepted range. Priced against the Vatican
    line (a flat per-person rate with no tier ceiling) so the assertion is
    about pax-parsing strictness, not about the Uffizi item's own 10-pax
    tier limit."""
    template, wb = _itinerary_with_pax_cell(itinerary, tmp_path, "200", row=9)

    lines, _skipped = parse_itinerary(wb, template)
    vatican_line = next(l for l in lines if l.row == 9)
    assert vatican_line.pax == 200
    assert vatican_line.pax_kind == "parsed"

    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(rate_sheet))
    vatican = next(p for p in priced if p.row == 9)
    assert vatican.status == "priced"
    assert vatican.total == Decimal("19000.00")   # 95 * 200, no PM supplement


@pytest.mark.parametrize("pax_value", NEEDS_QUOTE_PAX_CASES)
def test_ambiguous_pax_text_keeps_the_line_and_asks_for_confirmation(
        itinerary, rate_sheet, tmp_path, pax_value):
    """Every ambiguous or implausible pax cell -- ranges, open counts,
    extra words, out-of-bounds numbers, non-integers, and even a stray
    date -- must never vanish, never be priced, and never be silently
    filled in from default_pax even when one is configured. The line is
    kept and the reason names the actual cell content so the operator can
    see what confused the parser."""
    template, wb = _itinerary_with_pax_cell(itinerary, tmp_path, pax_value)
    template.default_pax = 4   # must not be used as a fallback here

    lines, _skipped = parse_itinerary(wb, template)
    assert any(l.row == 8 for l in lines), "line must not be silently dropped"
    uffizi_line = next(l for l in lines if l.row == 8)
    assert uffizi_line.pax_kind in ("unparseable", "out_of_range")

    priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))
    uffizi = next(p for p in priced if p.row == 8)
    assert uffizi.status == "needs_quote"
    assert uffizi.total is None

    quoted_text = pax_value.isoformat() if hasattr(pax_value, "isoformat") else pax_value
    assert quoted_text in uffizi.reason


def test_invalid_pax_never_ends_up_priced_with_a_missing_or_zero_total(
        itinerary, rate_sheet, tmp_path):
    """Blanket invariant across every case above (parsed, out-of-range,
    and unparseable alike): no returned line may ever be `status ==
    "priced"` while carrying `total is None` or `total == Decimal("0.00")`
    -- that combination is this module's own definition of a bug, whatever
    produced it."""
    all_cases = [c.values[0] for c in PARSED_PAX_CASES] + \
        [c.values[0] for c in NEEDS_QUOTE_PAX_CASES] + ["200"]

    for pax_value in all_cases:
        template, wb = _itinerary_with_pax_cell(itinerary, tmp_path, pax_value)
        template.default_pax = 4
        lines, _skipped = parse_itinerary(wb, template)
        priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))
        for line in priced:
            if line.status == "priced":
                assert line.total is not None, pax_value
                assert line.total != Decimal("0.00"), pax_value


# --- I6: the same invariant, varied along the RATE axis -------------------
#
# The test above only ever varies the PAX cell -- it was named as a
# "blanket invariant" but the RATE side was never actually exercised, so a
# mutation that broke either of the two rate-side guards could survive the
# whole suite undetected:
#   - pricing/patterns.py's `_lookup` raising NeedsQuote for a missing rate
#   - pricing/catalog.py's `resolve_rates` `continue`-ing past an
#     unparseable rate cell rather than resolving it to 0
# This sibling test widens the invariant to the RATE axis: a missing rate
# cell, an unparseable one ("on request"), an empty one, and a zero one all
# must come back needs_quote, never priced with a missing or non-positive
# total.

RATE_CELL_CASES = [
    pytest.param(None, id="missing"),
    pytest.param("on request", id="unparseable"),
    pytest.param("", id="empty"),
    pytest.param(0, id="zero"),
]


@pytest.mark.parametrize("rate_value", RATE_CELL_CASES)
def test_invalid_rate_never_ends_up_priced_with_a_missing_or_zero_total(
        itinerary, rate_sheet, tmp_path, rate_value):
    """Every one of these rate-cell shapes must come back needs_quote for
    the Vatican line (row 9, flat-rate, no modifiers, at Italy!D6) --
    never priced with a missing or non-positive total. Deliberately NOT
    the Uffizi line: it has a PM day_part supplement (Italy!F4, 40) that
    would mask a zeroed-out base rate (0 + 40 = 40, still positive), which
    would defeat the point of the "zero" case below."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["D6"] = rate_value
    mutated = tmp_path / "mutated_rate.xlsx"
    raw.save(mutated)

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(mutated))

    vatican = next(p for p in priced if p.row == 9)
    assert vatican.status == "needs_quote", rate_value
    assert vatican.total is None, rate_value


# The base-rate case above is deliberately chosen so an unresolvable rate
# collapses the whole total to 0 -- which means C1's own "total <= 0"
# guard in patterns.evaluate would ALSO catch a broken rate-side guard by
# coincidence, masking whether the rate-side guards themselves still
# work. An unresolvable MODIFIER cell is not masked this way: the base
# rate is still positive on its own, so if a missing/unparseable
# modifier were ever silently treated as "add 0" instead of raising, the
# line would come back priced at a plausible-looking but WRONG total
# (the supplement just silently missing) rather than 0 -- exactly the
# kind of error C1's guard cannot see. This is the test that actually
# isolates the two rate-side guards I6 is about. Zero is deliberately
# excluded here (unlike the base-rate cases above): a modifier that
# genuinely resolves to 0 (e.g. "no PM supplement this season") is not
# ambiguous the way a zero BASE rate is -- the line still charges a real,
# positive amount, so there's nothing to ask about.

MODIFIER_RATE_CELL_CASES = [
    pytest.param(None, id="missing"),
    pytest.param("on request", id="unparseable"),
    pytest.param("", id="empty"),
]


@pytest.mark.parametrize("rate_value", MODIFIER_RATE_CELL_CASES)
def test_invalid_modifier_rate_cell_never_silently_drops_the_modifier(
        itinerary, rate_sheet, tmp_path, rate_value):
    """The Uffizi line (row 8, PM, pax 4) has a positive base rate
    (Italy!D4) AND a PM day_part supplement (Italy!F4). If the supplement
    cell can't be resolved, the line must come back needs_quote -- never
    priced at base-rate-only as if the supplement were legitimately 0."""
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["F4"] = rate_value
    mutated = tmp_path / "mutated_modifier_rate.xlsx"
    raw.save(mutated)

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(mutated))

    uffizi = next(p for p in priced if p.row == 8)
    assert uffizi.status == "needs_quote", rate_value
    assert uffizi.total is None, rate_value


def test_default_pax_fills_in_when_there_is_no_pax_column(wide_itinerary, rate_sheet):
    wb = Workbook(wide_itinerary)
    template = propose_template(wb, "Itinerary")
    template.default_pax = 4
    lines, _skipped = parse_itinerary(wb, template)
    assert [l.pax for l in lines] == [4, 4]


def test_without_a_date_or_time_column_time_priced_items_need_a_quote(minimal_itinerary,
                                                                     rate_sheet):
    wb = Workbook(minimal_itinerary)
    template = propose_template(wb, "Itinerary")
    template.default_pax = 4
    lines, _skipped = parse_itinerary(wb, template)
    priced, _ = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(rate_sheet))
    uffizi = next(p for p in priced if p.raw_name.startswith("Uffizi"))
    assert uffizi.status == "needs_quote"
    assert "time of day" in uffizi.reason.lower()


def test_price_changes_are_surfaced(itinerary, rate_sheet, tmp_path):
    import openpyxl
    raw = openpyxl.load_workbook(rate_sheet)
    raw["Italy"]["D4"] = 310
    changed = tmp_path / "changed.xlsx"
    raw.save(changed)

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    _, changes = price_lines(lines, _catalog(), _aliases_with_uffizi(), Workbook(changed))
    assert any(c.cell == "Italy!D4" for c in changes)


# --- I1: except KeyError is narrowed to catalog.get, so a malformed
# catalog entry gets its own reason instead of a false "no item" message --

def test_genuinely_missing_catalog_item_still_says_so(itinerary, rate_sheet):
    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(rate_sheet),
                            resolutions={8: "NOT-A-REAL-ID"})
    line = next(p for p in priced if p.row == 8)
    assert line.status == "unmatched"
    assert line.reason == "Catalog has no item 'NOT-A-REAL-ID'."


def test_malformed_catalog_item_gets_a_distinct_reason_naming_what_is_missing(
        itinerary, rate_sheet):
    """A KeyError raised from inside resolve_rates/evaluate (not from
    catalog.get itself) used to be caught by the same `except KeyError`
    and mislabeled "Catalog has no item {item_id!r}" -- a plainly false
    reason, since the item IS in the catalog; it's just malformed (e.g.
    /map wrote a "pricing" dict with no "cell" key). She'd re-run /map and
    reproduce the exact same broken entry. This must get a distinct
    reason that actually names the problem."""
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-BROKEN", name="Broken Item", name_cell="Italy!A4",
        region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "flat"},   # malformed: missing "cell"
        source_rows=[4]))

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, catalog, Aliases(), Workbook(rate_sheet),
                            resolutions={8: "IT-BROKEN"})
    line = next(p for p in priced if p.row == 8)
    assert line.status == "unmatched"
    # Tightened to the full string, not a substring -- see
    # test_malformed_catalog_item_names_a_missing_sheet_cleanly below for
    # why a substring check let a mangled message through.
    assert line.reason == (
        "The catalog entry 'IT-BROKEN' looks malformed -- it is missing "
        "'cell'. Re-run /map for this item to fix it."
    )


def test_malformed_catalog_item_names_a_missing_sheet_cleanly(itinerary, rate_sheet):
    """IMPORTANT regression (code review): when the KeyError caught here
    comes from resolve_rates reading a rate cell whose SHEET has been
    renamed out from under it (not a missing dict key), the message used
    to read openpyxl's raw sentence back inside another sentence --
    "...it is missing 'Worksheet Italy does not exist.'". `rate_sheet`'s
    catalog item lives on the "Italy" sheet; `itinerary` (used here only
    as a stand-in rates workbook with no "Italy" sheet at all) reproduces
    the renamed/missing-sheet case without needing to actually rename
    anything."""
    catalog = Catalog()
    catalog.upsert(CatalogItem(
        id="IT-ROM-0001", name="Vatican Museums — Skip the Line",
        name_cell="Italy!A6", region="Italy", currency="EUR", unit="per_person",
        pricing={"pattern": "flat", "cell": "Italy!D6", "last_seen": "95"},
        source_rows=[6], row_fingerprint=fingerprint("Vatican Museums — Skip the Line")))

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    # `itinerary` (the workbook stood in for wb_rates here) has no "Italy"
    # sheet at all -- resolve_rates's own read of Italy!D6 is what raises.
    priced, _ = price_lines(lines, catalog, Aliases(), Workbook(itinerary),
                            resolutions={9: "IT-ROM-0001"})
    line = next(p for p in priced if p.row == 9)
    assert line.status == "unmatched"
    assert line.reason == (
        "The catalog entry 'IT-ROM-0001' looks malformed -- it is missing "
        "sheet 'Italy'. Re-run /map for this item to fix it."
    )


def test_no_score_threshold_backdoor_in_cost_module():
    """A close-but-unconfirmed match must stay unmatched even though it
    would rank first and score very high — there is no CONFIDENT_MATCH
    constant, and nothing in this module may compare `score` to a number."""
    import inspect

    import pricing.cost as cost_module

    assert not hasattr(cost_module, "CONFIDENT_MATCH")
    source = inspect.getsource(cost_module)
    assert ".score" not in source


def test_high_score_without_auto_accept_never_auto_prices(itinerary, rate_sheet, monkeypatch):
    """Behavioural companion to the text-based check above. That check only
    catches a literal `CONFIDENT_MATCH` name and a literal `.score`
    substring -- it is defeated by renaming the constant and reading the
    score via getattr(). This test is not: it stubs out candidates() to
    return a single candidate with a very high score and auto_accept=False,
    and asserts the line is never auto-priced regardless of how (or
    whether) the module happens to look at `score`."""
    import pricing.cost as cost_module

    stub_candidate = types.SimpleNamespace(
        item_id="IT-FLR-0032",
        name="Uffizi Gallery — Private Guided Tour",
        score=0.99,
        auto_accept=False,
    )
    monkeypatch.setattr(cost_module, "candidates", lambda *args, **kwargs: [stub_candidate])

    wb = Workbook(itinerary)
    lines, _skipped = parse_itinerary(wb, propose_template(wb, "Itinerary"))
    priced, _ = price_lines(lines, _catalog(), Aliases(), Workbook(rate_sheet))

    uffizi = next(p for p in priced if p.row == 8)
    assert uffizi.status == "unmatched"
    assert uffizi.total is None
    assert uffizi.item_id is None
