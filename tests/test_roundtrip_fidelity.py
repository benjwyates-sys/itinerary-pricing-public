"""Characterisation tests: what does an openpyxl save() actually destroy?

These assertions encode the *current measured* behaviour. If a future
openpyxl or Pillow changes it, these fail loudly and
docs/fidelity-findings.md / docs/fidelity.json need updating — which is
the point.
"""

import json
import zipfile
from pathlib import Path

import openpyxl

from pricing.workbook import CellRef, Workbook, copy_for_write


def _images(path):
    return len(openpyxl.load_workbook(path)["Itinerary"]._images)


def _cf_ranges(path):
    return len(list(openpyxl.load_workbook(path)["Itinerary"].conditional_formatting))


def _image_bytes(path):
    """Raw bytes of the embedded PNG, read straight from the zip.

    Bypasses openpyxl entirely so a silent re-encode (recompression,
    palette change, stripped metadata) on save can't hide behind
    openpyxl's own (possibly lossy) round trip of its own output.
    """
    with zipfile.ZipFile(path) as archive:
        return archive.read("xl/media/image1.png")


def test_copy_alone_preserves_everything(branded_itinerary, tmp_path):
    dst = copy_for_write(branded_itinerary, tmp_path / "copy.xlsx")
    assert _images(dst) == _images(branded_itinerary)
    assert _cf_ranges(dst) == _cf_ranges(branded_itinerary)
    assert _image_bytes(dst) == _image_bytes(branded_itinerary)


def test_save_roundtrip_preserves_conditional_formatting(branded_itinerary, tmp_path):
    dst = copy_for_write(branded_itinerary, tmp_path / "out.xlsx")
    wb = Workbook(dst)
    wb.write(CellRef.parse("Itinerary!E8"), 1140)
    wb.save()
    wb.close()
    assert _cf_ranges(dst) == _cf_ranges(branded_itinerary)


def test_save_roundtrip_records_image_behaviour(branded_itinerary, tmp_path):
    """Whatever happens to images, record it explicitly rather than assume.

    HISTORY: this project originally pinned dependencies to exactly
    `openpyxl` (no Pillow). Under that constraint, openpyxl's own reader
    (openpyxl/reader/drawings.py: `if not PILImage: # Pillow not installed,
    drop images`) silently dropped every embedded image on load — the image
    was invisible to openpyxl's in-memory model before any write or save
    ever ran, and `before` measured 0. See docs/fidelity-findings.md for
    that earlier finding, kept there as history.

    RULING: the project owner decided the agency logo on a client-facing
    costed itinerary matters more than dependency minimalism, so Pillow is
    now a declared runtime dependency (pyproject.toml).

    MEASURED with Pillow present (openpyxl 3.1.5, Pillow 12.3.0):
    `before` is 1 — openpyxl can now decode the embedded PNG on load. After
    copy_for_write -> Workbook load -> write -> save(), `after` is also 1.
    Count alone isn't proof the logo still looks right, though: a future
    openpyxl/Pillow that keeps the count at 1 but re-encodes the image
    (recompresses, strips metadata, changes the palette) would pass a
    count-only check while visibly degrading a client-facing document.
    So this test also compares the raw PNG bytes read directly out of the
    xlsx zip (bypassing openpyxl's own decoder) between the pristine
    fixture and the saved output: they are byte-identical.
    """
    dst = copy_for_write(branded_itinerary, tmp_path / "out.xlsx")
    before = _images(branded_itinerary)
    wb = Workbook(dst)
    wb.write(CellRef.parse("Itinerary!E8"), 1140)
    wb.save()
    wb.close()
    after = _images(dst)
    # Count assertions localise a failure where the image disappears
    # entirely (count drops to 0) rather than being silently altered.
    assert before == 1
    assert after == 1
    # Byte-identity assertion: the image content itself, not just its
    # presence, survives the round trip unchanged.
    assert _image_bytes(dst) == _image_bytes(branded_itinerary)


def test_fidelity_json_matches_measured_behaviour(branded_itinerary, tmp_path):
    """docs/fidelity.json is the machine-readable source of truth a later
    task greps^Wreads to decide whether to warn the operator about a lost
    logo. This test is what keeps it honest: it re-derives both booleans
    from the same round-trip this module's other tests exercise, and fails
    if docs/fidelity.json disagrees — so a behaviour change can't leave the
    JSON stale without also failing here.
    """
    fidelity_path = Path(__file__).resolve().parent.parent / "docs" / "fidelity.json"
    assert fidelity_path.exists(), "docs/fidelity.json is missing"

    data = json.loads(fidelity_path.read_text())

    dst = copy_for_write(branded_itinerary, tmp_path / "out.xlsx")
    wb = Workbook(dst)
    wb.write(CellRef.parse("Itinerary!E8"), 1140)
    wb.save()
    wb.close()

    measured_images_survive = (
        _images(dst) == _images(branded_itinerary) == 1
        and _image_bytes(dst) == _image_bytes(branded_itinerary)
    )
    measured_cf_survives = _cf_ranges(dst) == _cf_ranges(branded_itinerary)

    assert data["images_survive_roundtrip"] == measured_images_survive
    assert data["conditional_formatting_survives"] == measured_cf_survives
    assert data["openpyxl"] == openpyxl.__version__
    import PIL
    assert data["pillow"] == PIL.__version__
